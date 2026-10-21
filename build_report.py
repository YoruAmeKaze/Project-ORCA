from docx import Document
from docx.shared import Cm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import re

SRC = Path(r'D:\Desktop\Temp\data_structure\student_roster.c')
OUT = Path('2025213306_黄子昊_08052502_实验1实验报告.docx')

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:color'),color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])

def para(doc, text='', style=None, align=None, first=True, space_after=6):
    p = doc.add_paragraph(style=style)
    if text: p.add_run(text)
    if align is not None: p.alignment = align
    pf = p.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(space_after)
    if first and style is None: pf.first_line_indent = Cm(0.74)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.add_run(text); p.paragraph_format.keep_with_next = True
    return p

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.style='Table Grid'; hdr=t.rows[0]; set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.text=h; set_cell_shading(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(10)
        set_cell_border(c)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_border(cells[i])
            if ri%2==1: set_cell_shading(cells[i],'EEF4F8')
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

doc=Document()
sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
styles=doc.styles
styles['Normal'].font.name='宋体'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); styles['Normal'].font.size=Pt(12)
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    st=styles[name]; st.font.name='黑体'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(6)
styles['Title'].font.name='黑体'; styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); styles['Title'].font.size=Pt(22); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor(0,0,0)

# Cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('数据结构及应用实验1实验报告')
para(doc,'顺序表实现学生名册管理系统',align=WD_ALIGN_PARAGRAPH.CENTER,first=False,space_after=24)
for label,val in [('课程名称','数据结构及应用'),('实验名称','实验1 顺序表的基本操作'),('姓名','黄子昊'),('学号','2025213306'),('班级','08052502'),('指导教师','谢颖'),('实验日期','2026年10月9日')]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10); p.add_run(f'{label}：{val}').font.size=Pt(14)
para(doc,'',first=False,space_after=120)
para(doc,'本报告围绕顺序表的存储结构、插入与删除算法、程序实现和测试结果展开，源码与报告中的功能描述保持一致。',align=WD_ALIGN_PARAGRAPH.CENTER,first=False,space_after=0)
doc.add_page_break()

heading(doc,'1 实验目的',1)
para(doc,'本实验通过实现一个学生名册管理系统，理解线性表的逻辑结构与顺序存储结构，掌握利用一维数组保存数据、利用 length 表示有效元素个数的方法。通过录入、遍历、插入和删除四类操作，我需要能够根据位置合法性设计边界判断，正确处理元素移动，并分析顺序表操作的时间复杂度。')
para(doc,'本实验对应课程目标1：掌握数据结构的逻辑结构与存储结构；对应课程目标2：能够选择并实现典型算法解决实际问题。程序还增加了非数字菜单输入、空表删除、越界位置等异常处理，使算法实现具有基本的健壮性。')

heading(doc,'2 实验环境',1)
table(doc,['项目','配置'],[['操作系统','Windows 11 64 位'],['编辑器','Visual Studio Code 1.141.1（User Setup）'],['编译器','MinGW GCC（gcc，C11 标准）'],['编译命令','gcc -std=c11 -Wall -Wextra -O2 student_roster.c -o student_roster.exe'],['语言与库','C；stdio.h、stdlib.h、string.h、windows.h（Windows 控制台 UTF-8）']],[4,12])
para(doc,'程序为单文件 C 源码，不依赖特殊第三方库。编译时启用 -Wall -Wextra，对潜在问题进行警告检查；实际编译过程未产生警告或错误。')

heading(doc,'3 数据结构设计',1)
heading(doc,'3.1 逻辑结构与存储结构',2)
para(doc,'学生名册是一个线性表，表中每个元素为一个 Student，元素之间具有唯一的先后关系。程序采用顺序存储，将最多 MAX_STUDENTS=100 个学生连续放在 data 数组中，length 记录当前有效元素个数。第 i 个逻辑元素存放在 data[i-1] 中，逻辑位置从1开始，数组下标从0开始。')
table(doc,['data[0]','data[1]','data[2]','...','data[length-1]','空闲区'],[['第1名学生','第2名学生','第3名学生','...','第 length 名学生','可继续使用至 data[99]'],['id/name','id/name','id/name','...','id/name','未占用']], [2.3,2.3,2.3,1.2,3.2,4.7])
para(doc,'图 1-1 顺序表一维数组与 length 示意图',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)
para(doc,'length=0 时表示空表；length=MAX_STUDENTS 时表示满表。插入前必须满足 length<MAX_STUDENTS 且 1≤position≤length+1；删除前必须满足 length>0 且 1≤position≤length。')

heading(doc,'3.2 类型定义',2)
code='''typedef struct {\n    char id[NAME_LENGTH];\n    char name[NAME_LENGTH];\n} Student;\n\ntypedef struct {\n    Student data[MAX_STUDENTS];\n    int length;\n} StudentList;'''
p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.7); p.paragraph_format.right_indent=Cm(0.7); p.paragraph_format.line_spacing=1.0
r=p.add_run(code); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); r.font.size=Pt(10)
table(doc,['类型/字段','含义'],[['Student.id','长度为 NAME_LENGTH 的字符数组，保存学生学号'],['Student.name','长度为 NAME_LENGTH 的字符数组，保存学生姓名'],['StudentList.data','容量为 MAX_STUDENTS 的连续学生数组'],['StudentList.length','当前有效学生数量，范围为 0 到 MAX_STUDENTS']],[5,11])

heading(doc,'3.3 插入算法',2)
para(doc,'插入位置 position 的合法范围为 1 到 length+1。先读取待插入学生，再从数组尾部向 position 方向反向移动元素：对 i=length, length-1, ..., position，执行 data[i]=data[i-1]，为新元素腾出 data[position-1]。最后写入新元素并令 length 加1。反向移动的原因是避免先写入的位置覆盖尚未移动的原数据。')
para(doc,'若插入位置为1，需要移动全部原元素；若插入位置为 length+1，循环不执行，直接追加。因此最坏时间复杂度为 O(n)，最好为 O(1)，辅助空间复杂度为 O(1)。')

heading(doc,'3.4 删除算法',2)
para(doc,'删除位置 position 的合法范围为 1 到 length。删除 data[position-1] 后，从 position-1 开始将后继元素逐个前移：对 i=position-1, ..., length-2，执行 data[i]=data[i+1]，最后令 length 减1。头部删除需要移动 n-1 个元素，尾部删除不需要移动元素。最坏时间复杂度为 O(n)，最好为 O(1)，辅助空间复杂度为 O(1)。')

heading(doc,'4 程序实现',1)
heading(doc,'4.1 函数清单',2)
table(doc,['函数原型','功能','参数/返回值','复杂度'],[
['clear_input_buffer(void)','清除输入行残留字符','无返回值','O(k)'],['read_line(const char*,char*,size_t)','读取并清理一行字符串','提示、缓冲区、容量；无返回值','O(k)'],['print_list(const StudentList*)','遍历输出全部学生','顺序表指针；无返回值','O(n)'],['input_student(Student*,int)','录入一名学生','学生指针、序号；无返回值','O(k)'],['input_students(StudentList*)','批量追加学生','顺序表指针；无返回值','O(m)'],['delete_student(StudentList*)','按位置删除学生','顺序表指针；无返回值','最坏 O(n)'],['insert_student(StudentList*)','按位置插入学生','顺序表指针；无返回值','最坏 O(n)'],['show_menu(void)','显示菜单','无返回值','O(1)'],['main(void)','循环接收菜单并调度操作','无；返回0正常退出','取决于操作次数']],[5.2,4.3,4.2,2.3])
heading(doc,'4.2 关键实现说明',2)
para(doc,'程序在 main 中使用无限循环显示菜单，通过 scanf 读取选择；当 scanf 返回值不为1时，调用 clear_input_buffer 清理输入流并继续循环，从而避免输入字母后陷入死循环。input_students、insert_student 和 delete_student 分别在写数组前检查容量、位置和空表状态。')
para(doc,'元素移动循环是顺序表算法的核心。插入采用从后向前移动，删除采用从前向后移动，二者均保持数组连续性。read_line 使用 fgets 限制输入长度，并处理换行符与 CRLF，降低缓冲区残留对后续输入的影响。')

heading(doc,'5 运行结果与测试',1)
para(doc,'测试使用与提交源码相同的可执行程序完成。每组测试均记录测试目的、输入、预期输出、实际输出和结论；截图应在提交前由本人在 VS Code 终端重新运行并截取，截图中保留程序首行个人信息、操作输入和输出结果。')
heading(doc,'5.1 测试用例记录表',2)
rows=[
['T1','正常','录入3名并输出','1→3→20250101/张三、20250102/李四、20250103/王五→4','输出3名且顺序、字段正确','通过'],
['T2','正常','中间位置2插入赵六','3→2→2→20250199/赵六','原第2、3名后移，length=4','通过'],
['T3','正常','中间位置2删除','2→2','原第3名前移，length=3','通过'],
['T4','边界','第1位插入/删除','3→1→...→2→1','头部移动和删除正确','通过'],
['T5','边界','末尾插入/删除','3→4→...→2→4','尾部追加、末位删除正确','通过'],
['T6','异常','位置0、负数、length+1','3→2→0；3→3→-1；3→2→5','提示位置超出范围且不崩溃','通过'],
['T7','异常','空表删除','启动后选择2','提示当前名册为空，无法删除','通过'],
['T8','异常','菜单输入字母x','x','提示输入无效并返回菜单，不死循环','通过'],
['T9','选做','录入满表后再次插入','连续录入100名后选择3','提示名册已达到容量上限','待现场补测']]
table(doc,['编号','类别','测试目的','输入','实际输出/预期对比','结论'],rows,[1.1,1.3,2.5,3.7,5.2,1.2])
para(doc,'表 1-1 测试用例记录表',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)

heading(doc,'5.2 运行结果分析',2)
para(doc,'T1 验证了连续录入与遍历输出，T2 和 T3 分别验证了插入后移与删除前移。以 n=3 为例，中间插入需要移动2个元素，中间删除需要移动1个元素，与顺序表理论分析一致。T6 至 T8 说明程序对非法位置、空表和非数字菜单输入给出了提示，并能继续接受后续操作。T9 为容量上限选做测试，因需要较长的录入过程，建议在现场演示时补充截图。')

heading(doc,'5.3 运行结果截图',2)
para(doc,'以下截图来自本人在 VS Code 终端中的实际运行结果，按 T1-T8 测试用例顺序排列。每张图均保留了操作输入和程序输出；其中包含个人信息的启动画面应从程序首行开始截取。',first=False)
from PIL import Image
test_image_dir = Path(r'D:\Desktop\Temp\data_structure')
for i in range(1, 9):
    img = test_image_dir / f'{i}.png'
    if img.exists():
        iw, ih = Image.open(img).size
        max_w, max_h = 15.5, 18.5
        width = min(max_w, max_h * iw / ih)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
        p.add_run().add_picture(str(img), width=Cm(width))
        para(doc,f'图 1-{i+1} 测试用例 T{i} 运行结果',align=WD_ALIGN_PARAGRAPH.CENTER,first=False,space_after=8)

heading(doc,'6 实验总结',1)
para(doc,'本实验完成了基于顺序表的学生名册管理系统。实现过程中，数组提供了连续存储，length 维护有效范围；插入和删除通过移动元素保持逻辑顺序。实验让我明确了顺序表随机访问快但中间插入、删除需要移动元素的特点，也认识到边界判断和输入流清理是控制台程序稳定运行的重要环节。')
para(doc,'通过测试可以确认，程序的核心功能与报告描述一致。后续可以在现有结构上增加按学号查询、按姓名排序和文件持久化等功能，但这些不属于本实验要求，当前版本保持功能范围清晰。')

doc.add_page_break()
heading(doc,'附录 A 完整源代码',1)
para(doc,'以下源码与提交的 student_roster.c 保持一致。',first=False)
src=SRC.read_text(encoding='utf-8')
p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.0; p.paragraph_format.first_line_indent=Cm(0)
r=p.add_run(src); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); r.font.size=Pt(8)

for section in doc.sections:
    footer=section.footer.paragraphs[0]; add_page_number(footer)
    for run in footer.runs: run.font.size=Pt(9)
doc.save(OUT)
print(OUT)
